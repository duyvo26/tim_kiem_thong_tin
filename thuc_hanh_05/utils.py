"""
Thực hành 05 - Kỹ thuật Cluster Pruning
Tiền xử lý văn bản, Tokenization (VnCoreNLP), Trích xuất đa định dạng.
"""

import os
import re
import docx
from bs4 import BeautifulSoup
from pdfminer.high_level import extract_text
import py_vncorenlp
import numpy as np
import aspose.words as aw

# ─────────────────────────────────────────────
# Cấu hình đường dẫn
# ─────────────────────────────────────────────
DIR_ROOT = os.path.dirname(os.path.abspath(__file__))

# Sử dụng lại JAVA_HOME từ bài 03 nếu có sẵn trên máy (Đã kiểm tra tồn tại)
JAVA_HOME       = r"C:\Program Files\Java\jdk1.8.0_202"
VNCORENLP_DIR   = os.path.join(DIR_ROOT, "vncorenlp")
DATASET_DIR     = os.path.join(DIR_ROOT, "dataset _B5_1", "dataset _B5")
STOPWORDS_PATH  = os.path.join(DIR_ROOT, "vietnamese-stopwords.txt")

# Set JAVA env
os.environ["JAVA_HOME"] = JAVA_HOME
os.environ["PATH"] = os.path.join(JAVA_HOME, "bin") + os.pathsep + os.environ.get("PATH", "")

# ─────────────────────────────────────────────
# Load model VnCoreNLP (lazy – chỉ gọi 1 lần)
# ─────────────────────────────────────────────
_model = None

def get_model():
    global _model
    if _model is None:
        if not os.path.exists(VNCORENLP_DIR):
            print("Đang tải VnCoreNLP model...")
            py_vncorenlp.download_model(save_dir=VNCORENLP_DIR)
        print("Đang nạp VnCoreNLP model (wseg)...")
        _model = py_vncorenlp.VnCoreNLP(save_dir=VNCORENLP_DIR, annotators=["wseg"])
        print("VnCoreNLP model đã sẵn sàng.")
    return _model


# ─────────────────────────────────────────────
# Đọc file (Hỗ trợ đa định dạng)
# ─────────────────────────────────────────────
def read_docx(path: str) -> str:
    doc = docx.Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def read_pdf(path: str) -> str:
    return extract_text(path)


def read_html(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f, "html.parser")
    # Loại bỏ các tag rác
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript"]):
        tag.decompose()
    texts = []
    # Lấy nội dung chính
    for tag in soup.find_all(["p", "h1", "h2", "h3", "h4", "h5", "h6", "li"]):
        t = tag.get_text(separator=" ", strip=True)
        if t:
            texts.append(t)
    return "\n".join(texts)

def read_doc_legacy(path: str) -> str:
    
    try:
        # Load tài liệu .doc hoặc .docx
        doc = aw.Document(path)
        
        # Trích xuất toàn bộ text
        text = doc.get_text()
        
        # Loại bỏ các dòng thông báo bản quyền của bản Evaluation (dùng regex để xóa cả link)
        text = re.sub(r"Created with an evaluation copy of Aspose\.Words.*?license/", "", text, flags=re.DOTALL)
        
        # Làm sạch các ký tự đặc biệt Aspose thường thêm vào (như ký tự điều hướng hoặc mã lỗi)
        text = text.replace("Evaluation Only. Created with Aspose.Words. Copyright 2003-2024 Aspose Pty Ltd.", "")
        text = text.strip()
        
        return text
    except Exception as e:
        # print(f"  [Error] Aspose failed for {path}: {e}")
        return ""

def read_file(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return read_docx(path)
    elif ext == ".doc":
        return read_doc_legacy(path)
    elif ext == ".pdf":
        return read_pdf(path)
    elif ext in (".html", ".htm"):
        return read_html(path)
    else:
        raise ValueError(f"Không hỗ trợ định dạng file: {ext}")


# ─────────────────────────────────────────────
# Tiền xử lý (Stopwords & Tokenize)
# ─────────────────────────────────────────────
def load_stopwords(path: str = STOPWORDS_PATH):
    if not os.path.exists(path):
        return set()
    with open(path, "r", encoding="utf-8") as f:
        return set(line.strip() for line in f if line.strip())

def clean_text(text: str) -> str:
    """Chuẩn hóa văn bản cơ bản trước khi tokenize."""
    text = text.lower()
    # Loại bỏ dấu câu và ký tự lạ, giữ lại gạch dưới (cho VnCoreNLP)
    text = re.sub(r'[^\w\s]', ' ', text)
    # Loại bỏ khoảng trắng thừa
    text = re.sub(r'\s+', ' ', text).strip()
    return text

def get_tokens(text: str, stopwords: set, remove_stopwords: bool = True) -> list[str]:
    """
    Tokenize văn bản và trả về danh sách token đã xử lý.
    """
    if not text.strip():
        return []
    
    try:
        model = get_model()
        annotated = model.annotate_text(text)
        
        # model.annotate_text thường trả về dict {'0': [tokens], '1': [tokens]...} hoặc list
        sentences_raw = list(annotated.values()) if isinstance(annotated, dict) else annotated
        
        all_tokens = []
        for sent in sentences_raw:
            for w in sent:
                word = str(w.get("wordForm", "")).lower().strip()
                if not word or word == '_':
                    continue
                # Chỉ lọc token có chữ cái/số
                if not re.match(r"^[\w\u00C0-\u1EF9]+$", word):
                    continue
                if remove_stopwords and word in stopwords:
                    continue
                all_tokens.append(word)
        return all_tokens
    except Exception as e:
        print(f"  [WARN] Lỗi tokenize: {e}")
        return []

# ─────────────────────────────────────────────
# Xử lý toàn bộ tài liệu (D dùng cho Cluster Pruning)
# ─────────────────────────────────────────────
def load_all_documents(dataset_dir: str = DATASET_DIR, 
                        stopwords_path: str = STOPWORDS_PATH) -> dict:
    """
    Đọc tất cả tài liệu, trả về dict: { "filename": [token1, token2, ...] }
    """
    stopwords = load_stopwords(stopwords_path)
    documents = {}

    if not os.path.exists(dataset_dir):
        print(f"Lỗi: Không tìm thấy thư mục dataset tại {dataset_dir}")
        return {}

    files = sorted(f for f in os.listdir(dataset_dir)
                   if os.path.isfile(os.path.join(dataset_dir, f))
                   and not f.startswith("."))

    print(f"Đang đọc {len(files)} tài liệu...")
    for fname in files:
        fpath = os.path.join(dataset_dir, fname)
        try:
            text = read_file(fpath)
            tokens = get_tokens(text, stopwords)
            if tokens:
                documents[fname] = tokens
                print(f"  [OK] {fname}: {len(tokens)} tokens")
            else:
                print(f"  [Skip] {fname}: Không có nội dung")
        except Exception as e:
            print(f"  [Error] {fname}: {e}")

    print(f"\nTổng số tài liệu đã xử lý: {len(documents)}")
    return documents
